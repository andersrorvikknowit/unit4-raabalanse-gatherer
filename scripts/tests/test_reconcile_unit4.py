import importlib.util
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from decimal import Decimal
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


MODULE_PATH = Path(__file__).parents[1] / "reconcile_unit4.py"
SPEC = importlib.util.spec_from_file_location("reconcile_unit4", MODULE_PATH)
assert SPEC and SPEC.loader
reconcile = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = reconcile
SPEC.loader.exec_module(reconcile)


def grid(rows, formulas=()):
    return reconcile.Grid(
        "Sheet1",
        tuple(tuple(row) for row in rows),
        frozenset(formulas),
    )


def raw_grid(*account_rows, formulas=()):
    return grid(
        [
            ["Client:", "999999 Synthetic Example AS"],
            ["Period:", "209912"],
            [],
            ["Account", "Amount"],
            *account_rows,
        ],
        formulas,
    )


def periodised_grid(account, amounts, *, other_account=None, formulas=()):
    rows = [
        ["Client:", "999999 Synthetic Example AS"],
        ["Period:", "209912"],
        ["Konto:", account],
        [],
        ["Account", "Remaining Amount"],
    ]
    rows.extend([[account, amount] for amount in amounts])
    if other_account is not None:
        rows.append([other_account, "1,00"])
    total = sum((reconcile.require_cents(value) for value in amounts), Decimal("0.00"))
    if other_account is not None:
        total += Decimal("1.00")
    rows.append(["Total:", str(total)])
    return grid(rows, formulas)


class TextNormalizationTests(unittest.TestCase):
    def test_transliterates_norwegian_ae_and_o(self):
        self.assertEqual(reconcile.normalize_text("ÆØ æø"), "aeo aeo")


class MoneyParsingTests(unittest.TestCase):
    def test_parses_norwegian_and_english_formats(self):
        cases = {
            "3 289 422,29": Decimal("3289422.29"),
            "3\u00a0289\u202f422,29": Decimal("3289422.29"),
            "1.234.567,89": Decimal("1234567.89"),
            "1,234,567.89": Decimal("1234567.89"),
            "(1 234,50)": Decimal("-1234.50"),
            "1 234,50-": Decimal("-1234.50"),
            "NOK 1 234,50": Decimal("1234.50"),
        }
        for source, expected in cases.items():
            with self.subTest(source=source):
                self.assertEqual(reconcile.parse_money(source), expected)

    def test_rejects_subcent_precision(self):
        with self.assertRaisesRegex(reconcile.Blocked, "subcent"):
            reconcile.require_cents(Decimal("1.001"))

    def test_accepts_excel_float_noise_but_not_real_subcent(self):
        self.assertEqual(reconcile.require_cents(402081.66000000003), Decimal("402081.66"))
        with self.assertRaisesRegex(reconcile.Blocked, "subcent"):
            reconcile.require_cents(100.001)

    def test_rejects_malformed_space_grouping(self):
        with self.assertRaisesRegex(reconcile.Blocked, "ambiguous"):
            reconcile.parse_money("1234 56,78")

    def test_rejects_single_separator_with_three_trailing_digits(self):
        for source in ("1,000", "1.234", "1234,567", "1234.567"):
            with self.subTest(source=source):
                with self.assertRaisesRegex(reconcile.Blocked, "ambiguous"):
                    reconcile.parse_money(source)


class IdentityTests(unittest.TestCase):
    def test_accepts_matching_labeled_identity(self):
        reconcile.validate_identity(
            (grid([["Company: 999999 Synthetic Example AS"], ["Periode", 209912]]),),
            "999999",
            "209912",
        )

    def test_rejects_conflicting_period(self):
        with self.assertRaisesRegex(reconcile.Blocked, "period"):
            reconcile.validate_identity(
                (grid([["Client:", "999999 Synthetic Example AS"], ["Period:", "209911"]]),),
                "999999",
                "209912",
            )

    def test_rejects_conflicting_leading_company_even_when_name_contains_request(self):
        with self.assertRaisesRegex(reconcile.Blocked, "conflicting_company"):
            reconcile.validate_identity(
                (
                    grid(
                        [
                            ["Client:", "123 Company 999999 AS"],
                            ["Period:", "209912"],
                        ]
                    ),
                ),
                "999999",
                "209912",
            )

    def test_accepts_company_token_when_no_leading_code_exists(self):
        reconcile.validate_identity(
            (
                grid(
                    [
                        ["Client:", "Synthetic company 999999"],
                        ["Period:", "209912"],
                    ]
                ),
            ),
            "999999",
            "209912",
        )


class ReconciliationTests(unittest.TestCase):
    def reconcile(
        self,
        raw,
        report_1510,
        report_1560,
        *,
        supplier_reports=None,
    ):
        raw_rows = list(raw.rows)
        existing_accounts = {
            str(row[0]).strip()
            for row in raw_rows
            if row and row[0] is not None
        }
        for account in reconcile.SUPPLIER_ACCOUNTS:
            if account not in existing_accounts:
                raw_rows.append((account, "0,00"))
        raw = reconcile.Grid(raw.name, tuple(raw_rows), raw.formula_cells)
        reports = {
            "1510": (report_1510,),
            "1560": (report_1560,),
            **{
                account: (periodised_grid(account, ["0,00"]),)
                for account in reconcile.SUPPLIER_ACCOUNTS
            },
        }
        if supplier_reports:
            reports.update(
                {account: (report,) for account, report in supplier_reports.items()}
            )
        return reconcile.reconcile_grids(
            (raw,),
            reports,
            company="999999",
            period="209912",
        )

    def test_matches_customer_accounts_with_zero_supplier_balances(self):
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            periodised_grid("1510", ["40,00", "60,00"]),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["overall_status"], "MATCH")
        self.assertEqual(
            [item["status"] for item in result["accounts"]],
            ["MATCH"] * 5,
        )

    def test_matches_supplier_accounts_from_separate_reports(self):
        result = self.reconcile(
            raw_grid(
                [1510, "100,00"],
                [1560, "25,00"],
                [2440, "40,00"],
                [2444, "44,00"],
                [2460, "60,00"],
            ),
            periodised_grid("1510", ["100,00"]),
            periodised_grid("1560", ["25,00"]),
            supplier_reports={
                "2440": periodised_grid("2440", ["10,00", "30,00"]),
                "2444": periodised_grid("2444", ["44,00"]),
                "2460": periodised_grid("2460", ["25,00", "35,00"]),
            },
        )
        self.assertEqual(result["overall_status"], "MATCH")
        self.assertEqual(
            [item["account"] for item in result["accounts"]],
            ["1510", "1560", "2440", "2444", "2460"],
        )
        self.assertEqual(
            [item["status"] for item in result["accounts"]],
            ["MATCH"] * 5,
        )

    def test_empty_supplier_is_zero_only_with_attested_order(self):
        raw = raw_grid(
            [1510, "0,00"], [1560, "0,00"], [2440, "0,00"],
            [2444, "0,00"], [2460, "0,00"],
        )
        reports = {
            account: (periodised_grid(account, []),)
            for account in reconcile.SUPPORTED_ACCOUNTS
            if account != "2444"
        }
        result = reconcile.reconcile_grids(
            (raw,), reports, company="999999", period="209912",
            empty_supplier_orders={"2444": "15192"},
        )
        self.assertEqual(result["overall_status"], "MATCH")
        self.assertEqual(result["accounts"][3]["periodised_sum"], "0.00")
        self.assertEqual(result["accounts"][3]["empty_result_order"], "15192")

        missing_raw = raw_grid(
            [1510, "0,00"], [1560, "0,00"], [2440, "0,00"], [2460, "0,00"],
        )
        result = reconcile.reconcile_grids(
            (missing_raw,), reports, company="999999", period="209912",
            empty_supplier_orders={"2444": "15192"},
        )
        self.assertEqual(result["accounts"][3]["status"], "MATCH")
        self.assertTrue(result["accounts"][3]["raw_row_absent_treated_as_zero"])

        nonzero_raw = raw_grid(
            [1510, "0,00"], [1560, "0,00"], [2440, "0,00"],
            [2444, "44,00"], [2460, "0,00"],
        )
        result = reconcile.reconcile_grids(
            (nonzero_raw,), reports, company="999999", period="209912",
            empty_supplier_orders={"2444": "15192"},
        )
        self.assertEqual(result["accounts"][3]["status"], "AVVIK")
        self.assertEqual(result["accounts"][3]["difference"], "-44.00")

    def test_matches_exact_norwegian_amount_headers(self):
        raw = grid(
            [
                ["Klient:", "999999 Synthetic Example AS"],
                ["Periode:", "209912"],
                [],
                ["Konto", "Beløp"],
                [1510, "100,00"],
                [1560, "25,00"],
            ]
        )

        def norwegian_periodised(account, amount):
            return grid(
                [
                    ["Klient:", "999999 Synthetic Example AS"],
                    ["Periode:", "209912"],
                    ["Konto:", account],
                    [],
                    ["Konto", "Gjenstående beløp"],
                    [account, amount],
                    ["Sum:", amount],
                ]
            )

        result = self.reconcile(
            raw,
            norwegian_periodised("1510", "100,00"),
            norwegian_periodised("1560", "25,00"),
        )

        self.assertEqual(result["overall_status"], "MATCH")
        self.assertEqual(
            [item["status"] for item in result["accounts"]],
            ["MATCH"] * 5,
        )

    def test_reports_signed_deviation(self):
        result = self.reconcile(
            raw_grid([1510, "99,00"], [1560, "25,00"]),
            periodised_grid("1510", ["40,00", "60,00"]),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["overall_status"], "AVVIK")
        self.assertEqual(result["accounts"][0]["difference"], "1.00")

    def test_rejects_mixed_account_periodised_report(self):
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            periodised_grid("1510", ["100,00"], other_account="1560"),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["overall_status"], "BLOCKED")
        self.assertEqual(
            result["accounts"][0]["reason"],
            "mixed_accounts_in_periodised_report",
        )

    def test_accepts_descriptive_account_parameter_for_empty_report(self):
        report_1510 = grid(
            [
                ["Client:", "999999 Synthetic Example AS"],
                ["Period:", "209912"],
                ["Konto:", "1510 - Kundefordringer"],
                [],
                ["Account", "Remaining Amount"],
                ["Total:", "0,00"],
            ]
        )
        result = self.reconcile(
            raw_grid([1510, "0,00"], [1560, "25,00"]),
            report_1510,
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["accounts"][0]["status"], "MATCH")

    def test_rejects_multiple_accounts_in_parameter(self):
        report_1510 = grid(
            [
                ["Client:", "999999 Synthetic Example AS"],
                ["Period:", "209912"],
                ["Konto:", "1510 / 1560"],
                [],
                ["Account", "Remaining Amount"],
                ["Total:", "0,00"],
            ]
        )
        result = self.reconcile(
            raw_grid([1510, "0,00"], [1560, "25,00"]),
            report_1510,
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(
            result["accounts"][0]["reason"],
            "ambiguous_periodised_account_parameter",
        )

    def test_rejects_wildcard_account_parameter(self):
        report_1510 = grid(
            [
                ["Client:", "999999 Synthetic Example AS"],
                ["Period:", "209912"],
                ["Konto:", "*"],
                [],
                ["Account", "Remaining Amount"],
                [1510, "100,00"],
                ["Total:", "100,00"],
            ]
        )
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            report_1510,
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(
            result["accounts"][0]["reason"],
            "wildcard_periodised_account_parameter",
        )

    def test_rejects_duplicate_and_missing_raw_rows_independently(self):
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1510, "100,00"]),
            periodised_grid("1510", ["100,00"]),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["accounts"][0]["reason"], "duplicate_raw_account_rows")
        self.assertEqual(result["accounts"][1]["reason"], "raw_account_row_missing")

    def test_rejects_relevant_formula_cells(self):
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"], formulas={(4, 1)}),
            periodised_grid("1510", ["100,00"]),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["overall_status"], "BLOCKED")
        self.assertEqual(result["accounts"][0]["reason"], "formula_in_raw_amount")
        self.assertEqual(result["accounts"][1]["status"], "MATCH")

    def test_rejects_formula_total_even_without_a_cached_value(self):
        report = periodised_grid("1510", ["100,00"])
        changed_rows = list(report.rows)
        changed_rows[-1] = ("Total:", None)
        report = grid(changed_rows, formulas={(len(changed_rows) - 1, 1)})
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            report,
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(
            result["accounts"][0]["reason"], "formula_in_periodised_total"
        )

    def test_rejects_detail_total_mismatch(self):
        report = periodised_grid("1510", ["100,00"])
        changed_rows = list(report.rows)
        changed_rows[-1] = ("Total:", "99,00")
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            grid(changed_rows),
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(
            result["accounts"][0]["reason"],
            "periodised_total_does_not_match_details",
        )

    def test_unrelated_total_text_does_not_hide_detail_row(self):
        report_1510 = grid(
            [
                ["Client:", "999999 Synthetic Example AS"],
                ["Period:", "209912"],
                ["Konto:", "1510"],
                [],
                ["Customer", "Account", "Remaining Amount"],
                ["Total", "1510", "100,00"],
                [None, "Total:", "100,00"],
            ]
        )
        result = self.reconcile(
            raw_grid([1510, "100,00"], [1560, "25,00"]),
            report_1510,
            periodised_grid("1560", ["25,00"]),
        )
        self.assertEqual(result["accounts"][0]["status"], "MATCH")


@unittest.skipUnless(openpyxl is not None, "openpyxl is required for loader integration")
class LoaderCliTests(unittest.TestCase):
    @staticmethod
    def write_workbook(path, rows):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    def test_cli_loads_xlsx_and_returns_match_then_deviation(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            raw_path = directory / "raw.xlsx"
            values = {
                "1510": 100,
                "1560": 25,
                "2440": 40,
                "2444": 44,
                "2460": 60,
            }
            report_paths = {
                account: directory / f"{account}.xlsx"
                for account in reconcile.SUPPORTED_ACCOUNTS
            }
            self.write_workbook(
                raw_path,
                [
                    ["Client:", "999999 Synthetic Example AS"],
                    ["Period:", "209912"],
                    [],
                    ["Account", "Amount"],
                    *[[int(account), amount] for account, amount in values.items()],
                ],
            )
            for account, amount in values.items():
                self.write_workbook(
                    report_paths[account],
                    [
                        ["Client:", "999999 Synthetic Example AS"],
                        ["Period:", "209912"],
                        ["Konto:", int(account)],
                        [],
                        ["Account", "Remaining Amount"],
                        [int(account), amount],
                        ["Total:", amount],
                    ],
                )
            arguments = [
                "--raw-balance",
                str(raw_path),
            ]
            for account in reconcile.SUPPORTED_ACCOUNTS:
                arguments.extend(
                    ["--periodised", f"{account}={report_paths[account]}"]
                )
            arguments.extend(["--company", "999999", "--period", "209912"])

            output = io.StringIO()
            with redirect_stdout(output):
                exit_code = reconcile.main(arguments)
            payload = json.loads(output.getvalue())
            self.assertEqual(exit_code, 0)
            self.assertEqual(payload["overall_status"], "MATCH")
            self.assertNotIn("Synthetic Example AS", output.getvalue())

            workbook = openpyxl.load_workbook(raw_path)
            workbook.active.cell(row=5, column=2).value = 99
            workbook.save(raw_path)
            workbook.close()

            output = io.StringIO()
            with redirect_stdout(output):
                exit_code = reconcile.main(arguments)
            payload = json.loads(output.getvalue())
            self.assertEqual(exit_code, 2)
            self.assertEqual(payload["overall_status"], "AVVIK")
            self.assertEqual(payload["accounts"][0]["difference"], "1.00")

    def test_cli_preserves_valid_account_when_other_workbook_cannot_load(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            raw_path = directory / "raw.xlsx"
            values = {
                "1510": 100,
                "1560": 25,
                "2440": 40,
                "2444": 44,
                "2460": 60,
            }
            report_paths = {
                account: directory / f"{account}.xlsx"
                for account in reconcile.SUPPORTED_ACCOUNTS
            }
            self.write_workbook(
                raw_path,
                [
                    ["Client:", "999999 Synthetic Example AS"],
                    ["Period:", "209912"],
                    [],
                    ["Account", "Amount"],
                    *[[int(account), amount] for account, amount in values.items()],
                ],
            )
            for account, amount in values.items():
                if account == "1560":
                    report_paths[account].write_bytes(b"not an xlsx workbook")
                    continue
                self.write_workbook(
                    report_paths[account],
                    [
                        ["Client:", "999999 Synthetic Example AS"],
                        ["Period:", "209912"],
                        ["Konto:", int(account)],
                        [],
                        ["Account", "Remaining Amount"],
                        [int(account), amount],
                        ["Total:", amount],
                    ],
                )

            arguments = ["--raw-balance", str(raw_path)]
            for account in reconcile.SUPPORTED_ACCOUNTS:
                arguments.extend(
                    ["--periodised", f"{account}={report_paths[account]}"]
                )
            arguments.extend(["--company", "999999", "--period", "209912"])

            output = io.StringIO()
            with redirect_stdout(output):
                exit_code = reconcile.main(arguments)

            payload = json.loads(output.getvalue())
            self.assertEqual(exit_code, 3)
            self.assertEqual(payload["overall_status"], "BLOCKED")
            self.assertEqual(payload["accounts"][0]["status"], "MATCH")
            self.assertEqual(payload["accounts"][1]["status"], "BLOCKED")
            self.assertEqual(
                payload["accounts"][1]["reason"], "workbook_could_not_be_opened"
            )
            self.assertEqual(
                [item["status"] for item in payload["accounts"][2:]],
                ["MATCH", "MATCH", "MATCH"],
            )
            self.assertNotIn("not an xlsx workbook", output.getvalue())


if __name__ == "__main__":
    unittest.main()
