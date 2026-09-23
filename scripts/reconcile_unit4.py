#!/usr/bin/env python3
"""Deterministically reconcile Unit4 Råbalanse and customer-balance workbooks.

The program is intentionally read-only.  It prints a single JSON document and
never emits customer-level rows or customer names.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence


SUPPORTED_ACCOUNTS = ("1510", "1560")
CENT = Decimal("0.01")

COMPANY_LABELS = {"client", "company", "selskap", "firma", "klient"}
PERIOD_LABELS = {"period", "periode"}
ACCOUNT_HEADERS = {
    "account",
    "account no",
    "account number",
    "konto",
    "konto nr",
    "kontonr",
    "kontonummer",
}
RAW_AMOUNT_HEADERS = {
    "amount",
    "amount in local currency",
    "balance",
    "closing balance",
    "belop",
    "saldo",
    "utgaende saldo",
}
REMAINING_AMOUNT_HEADERS = {
    "remaining amount",
    "outstanding amount",
    "open amount",
    "restbelop",
    "gjenstaende belop",
}
TOTAL_LABELS = {"sum", "total", "totalt"}
NORWEGIAN_TRANSLITERATION = str.maketrans(
    {
        "Æ": "AE",
        "Ø": "O",
        "æ": "ae",
        "ø": "o",
    }
)


class ReconciliationError(Exception):
    """Base exception with a machine-readable, privacy-safe reason."""

    def __init__(self, reason: str, *, location: str | None = None):
        super().__init__(reason)
        self.reason = reason
        self.location = location


class Blocked(ReconciliationError):
    pass


class InputError(ReconciliationError):
    pass


class JsonArgumentParser(argparse.ArgumentParser):
    def error(self, message: str) -> None:
        raise InputError("invalid_cli_arguments")


@dataclass(frozen=True)
class Grid:
    """A cached-value worksheet plus coordinates that contain formulas."""

    name: str
    rows: tuple[tuple[Any, ...], ...]
    formula_cells: frozenset[tuple[int, int]] = frozenset()

    def value(self, row: int, column: int) -> Any:
        if row < 0 or row >= len(self.rows):
            return None
        values = self.rows[row]
        if column < 0 or column >= len(values):
            return None
        return values[column]

    def is_formula(self, row: int, column: int) -> bool:
        return (row, column) in self.formula_cells


@dataclass(frozen=True)
class Table:
    grid: Grid
    header_row: int
    account_column: int
    amount_column: int


@dataclass(frozen=True)
class LabeledValue:
    value: Any
    grid: Grid
    row: int
    column: int


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).casefold()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    # NFKD decomposes letters such as å, but not Norwegian æ or ø. Translate
    # those explicitly so documented Unit4 headers such as "Beløp" survive
    # the subsequent ASCII-only semantic normalization.
    text = str(value).translate(NORWEGIAN_TRANSLITERATION)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold().strip()
    text = re.sub(r"[^0-9a-z]+", " ", text)
    return " ".join(text.split())


def _cell_location(grid: Grid, row: int, column: int) -> str:
    number = column + 1
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    # Do not echo arbitrary sheet names: a workbook author can put customer
    # data there.  The coordinate is enough to diagnose a single chosen table.
    return f"cell:{letters}{row + 1}"


def _require_literal(grid: Grid, row: int, column: int, reason: str) -> None:
    if grid.is_formula(row, column):
        raise Blocked(reason, location=_cell_location(grid, row, column))


def _validate_thousands(integer_part: str, separator: str) -> str:
    groups = integer_part.split(separator)
    if not groups or not groups[0].isdigit() or not 1 <= len(groups[0]) <= 3:
        raise Blocked("ambiguous_monetary_value")
    if any(not group.isdigit() or len(group) != 3 for group in groups[1:]):
        raise Blocked("ambiguous_monetary_value")
    return "".join(groups)


def parse_money(value: Any) -> Decimal:
    """Parse numeric, Norwegian, or English monetary values without floats."""

    if value is None or isinstance(value, bool):
        raise Blocked("missing_or_invalid_monetary_value")
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise Blocked("missing_or_invalid_monetary_value")
        number = Decimal(str(value))
    else:
        text = unicodedata.normalize("NFKC", str(value)).strip()
        if not text:
            raise Blocked("missing_or_invalid_monetary_value")
        text = re.sub(r"(?i)\b(?:nok|kr)\b", "", text).strip()

        negative = False
        if text.startswith("(") and text.endswith(")"):
            negative = True
            text = text[1:-1].strip()
        elif text.startswith("(") or text.endswith(")"):
            raise Blocked("ambiguous_monetary_value")

        if text.endswith("-"):
            if negative:
                raise Blocked("ambiguous_monetary_value")
            negative = True
            text = text[:-1].strip()
        if text.startswith(("+", "-")):
            if text[0] == "-":
                if negative:
                    raise Blocked("ambiguous_monetary_value")
                negative = True
            text = text[1:].strip()

        spaced_text = re.sub(r"[\s\u00a0\u202f\u2009]+", " ", text).strip()
        if " " in spaced_text:
            space_groups = spaced_text.split(" ")
            if (
                not re.fullmatch(r"\d{1,3}", space_groups[0])
                or any(
                    not re.fullmatch(r"\d{3}", group)
                    for group in space_groups[1:-1]
                )
                or not re.fullmatch(r"\d{3}(?:[.,]\d+)?", space_groups[-1])
            ):
                raise Blocked("ambiguous_monetary_value")
        text = spaced_text.replace(" ", "")
        if not text or not re.fullmatch(r"\d[\d.,']*", text):
            raise Blocked("missing_or_invalid_monetary_value")

        if "'" in text:
            integer_candidate, dot, fractional_candidate = text.partition(".")
            if "," in integer_candidate or "'" in fractional_candidate:
                raise Blocked("ambiguous_monetary_value")
            integer_candidate = _validate_thousands(integer_candidate, "'")
            text = integer_candidate + (dot + fractional_candidate if dot else "")

        comma_position = text.rfind(",")
        dot_position = text.rfind(".")
        decimal_separator: str | None = None

        if comma_position >= 0 and dot_position >= 0:
            decimal_separator = "," if comma_position > dot_position else "."
            thousands_separator = "." if decimal_separator == "," else ","
            integer_part, fractional_part = text.rsplit(decimal_separator, 1)
            if decimal_separator in integer_part or not fractional_part.isdigit():
                raise Blocked("ambiguous_monetary_value")
            integer_digits = _validate_thousands(integer_part, thousands_separator)
            canonical = f"{integer_digits}.{fractional_part}"
        elif comma_position >= 0 or dot_position >= 0:
            separator = "," if comma_position >= 0 else "."
            groups = text.split(separator)
            if any(not group.isdigit() for group in groups):
                raise Blocked("ambiguous_monetary_value")
            if len(groups) == 2:
                # A single separator followed by three digits is inherently
                # locale-ambiguous (1,000 can mean one or one thousand).  Do
                # not silently scale it in either direction.
                if len(groups[1]) == 3:
                    raise Blocked("ambiguous_monetary_value")
                decimal_separator = separator
                canonical = f"{groups[0]}.{groups[1]}"
            elif all(len(group) == 3 for group in groups[1:]):
                canonical = "".join(groups)
            elif len(groups[-1]) <= 2 and all(
                len(group) == 3 for group in groups[1:-1]
            ):
                decimal_separator = separator
                canonical = f"{''.join(groups[:-1])}.{groups[-1]}"
            else:
                raise Blocked("ambiguous_monetary_value")
        else:
            canonical = text

        try:
            number = Decimal(canonical)
        except InvalidOperation as error:
            raise Blocked("missing_or_invalid_monetary_value") from error
        if negative:
            number = -number

    if not number.is_finite():
        raise Blocked("missing_or_invalid_monetary_value")
    return number


def require_cents(value: Any, *, location: str | None = None) -> Decimal:
    number = parse_money(value)
    try:
        cents = number.quantize(CENT)
    except InvalidOperation as error:
        raise Blocked("missing_or_invalid_monetary_value", location=location) from error
    if cents != number:
        raise Blocked("subcent_precision_is_not_supported", location=location)
    return cents


def normalize_account(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return None
        if number == number.to_integral_value():
            return str(number.quantize(Decimal("1")))
        return None
    text = unicodedata.normalize("NFKC", str(value)).strip()
    match = re.fullmatch(r"(\d+)(?:[.,]0+)?", text)
    return match.group(1) if match else None


def parameter_account_code(value: Any) -> str | None:
    """Read one leading account code while allowing only descriptive suffixes."""

    exact = normalize_account(value)
    if exact is not None:
        return exact
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).strip()
    codes = re.findall(r"(?<!\d)(\d{4,10})(?!\d)", text)
    if len(codes) != 1:
        return None
    if not re.match(rf"^\s*{re.escape(codes[0])}(?:\s|[-–—:])", text):
        return None
    return codes[0]


def _labeled_values(grids: Sequence[Grid], labels: set[str]) -> list[LabeledValue]:
    found: list[LabeledValue] = []
    for grid in grids:
        for row_index, row in enumerate(grid.rows):
            for column_index, raw in enumerate(row):
                normalized = normalize_text(raw)
                if not normalized:
                    continue
                for label in labels:
                    if normalized == label:
                        if isinstance(raw, str) and "*" in raw:
                            found.append(LabeledValue("*", grid, row_index, column_index))
                            continue
                        for next_column in range(column_index + 1, min(len(row), column_index + 5)):
                            candidate = row[next_column]
                            if candidate is not None and (
                                not isinstance(candidate, str) or candidate.strip()
                            ):
                                found.append(
                                    LabeledValue(candidate, grid, row_index, next_column)
                                )
                                break
                    elif normalized.startswith(label + " "):
                        remainder = normalized[len(label) + 1 :]
                        if remainder:
                            found.append(LabeledValue(remainder, grid, row_index, column_index))
    return found


def validate_identity(grids: Sequence[Grid], company: str, period: str) -> None:
    company_values = _labeled_values(grids, COMPANY_LABELS)
    matching_company = False
    conflicting_company = False
    company_pattern = re.compile(rf"(?<!\d){re.escape(company)}(?!\d)")
    for labeled in company_values:
        _require_literal(
            labeled.grid,
            labeled.row,
            labeled.column,
            "formula_in_company_identity",
        )
        text = normalize_text(labeled.value)
        leading_code = re.match(r"^(\d{2,10})(?:\s|$)", text)
        if leading_code:
            if leading_code.group(1) == company:
                matching_company = True
            else:
                conflicting_company = True
        elif company_pattern.search(text):
            # Some exports put the client code after the company name. Use that
            # weaker token match only when no authoritative leading code exists.
            matching_company = True
    if conflicting_company:
        raise Blocked("conflicting_company_identity")
    if not matching_company:
        raise Blocked("company_identity_not_found")

    period_values = _labeled_values(grids, PERIOD_LABELS)
    matching_period = False
    conflicting_period = False
    for labeled in period_values:
        _require_literal(
            labeled.grid,
            labeled.row,
            labeled.column,
            "formula_in_period_identity",
        )
        periods = re.findall(r"(?<!\d)(\d{6})(?!\d)", normalize_text(labeled.value))
        matching_period = matching_period or period in periods
        conflicting_period = conflicting_period or any(item != period for item in periods)
    if not matching_period:
        raise Blocked("period_identity_not_found")
    if conflicting_period:
        raise Blocked("conflicting_period_identity")


def find_table(
    grids: Sequence[Grid],
    amount_headers: set[str],
    *,
    report_kind: str,
) -> Table:
    candidates: list[Table] = []
    for grid in grids:
        for row_index, row in enumerate(grid.rows):
            normalized = [normalize_text(value) for value in row]
            account_columns = [
                index for index, value in enumerate(normalized) if value in ACCOUNT_HEADERS
            ]
            amount_columns = [
                index for index, value in enumerate(normalized) if value in amount_headers
            ]
            for account_column in account_columns:
                for amount_column in amount_columns:
                    if account_column == amount_column:
                        continue
                    _require_literal(
                        grid,
                        row_index,
                        account_column,
                        f"formula_in_{report_kind}_header",
                    )
                    _require_literal(
                        grid,
                        row_index,
                        amount_column,
                        f"formula_in_{report_kind}_header",
                    )
                    candidates.append(
                        Table(grid, row_index, account_column, amount_column)
                    )
    if not candidates:
        raise Blocked(f"{report_kind}_table_not_found")
    if len(candidates) != 1:
        raise Blocked(f"ambiguous_{report_kind}_table")
    return candidates[0]


def _is_total_row(table: Table, row_index: int) -> bool:
    """Recognize a total only in the semantic account/row-label column."""

    return (
        normalize_text(table.grid.value(row_index, table.account_column))
        in TOTAL_LABELS
    )


def extract_raw_value(table: Table, account: str) -> Decimal:
    matches: list[tuple[int, Any]] = []
    for row_index in range(table.header_row + 1, len(table.grid.rows)):
        if _is_total_row(table, row_index):
            continue
        raw_account = table.grid.value(row_index, table.account_column)
        if normalize_account(raw_account) == account:
            _require_literal(
                table.grid,
                row_index,
                table.account_column,
                "formula_in_raw_account",
            )
            matches.append((row_index, table.grid.value(row_index, table.amount_column)))
    if not matches:
        raise Blocked("raw_account_row_missing")
    if len(matches) != 1:
        raise Blocked("duplicate_raw_account_rows")
    row_index, value = matches[0]
    _require_literal(
        table.grid,
        row_index,
        table.amount_column,
        "formula_in_raw_amount",
    )
    return require_cents(
        value,
        location=_cell_location(table.grid, row_index, table.amount_column),
    )


def _account_parameter_values(table: Table) -> list[LabeledValue]:
    prefix = Grid(
        table.grid.name,
        table.grid.rows[: table.header_row],
        frozenset(
            coordinate
            for coordinate in table.grid.formula_cells
            if coordinate[0] < table.header_row
        ),
    )
    return _labeled_values((prefix,), ACCOUNT_HEADERS)


def extract_periodised_sum(table: Table, account: str) -> tuple[Decimal, Decimal | None]:
    parameters = _account_parameter_values(table)
    exact_parameter = False
    for labeled in parameters:
        _require_literal(
            labeled.grid,
            labeled.row,
            labeled.column,
            "formula_in_periodised_account_parameter",
        )
        normalized = parameter_account_code(labeled.value)
        if normalized == account:
            exact_parameter = True
        elif str(labeled.value).strip() == "*":
            raise Blocked("wildcard_periodised_account_parameter")
        elif normalized is not None:
            raise Blocked("conflicting_periodised_account_parameter")
        elif re.search(r"(?<!\d)\d{4,10}(?!\d)", str(labeled.value)):
            raise Blocked("ambiguous_periodised_account_parameter")

    amounts: list[Decimal] = []
    totals: list[Decimal] = []
    for row_index in range(table.header_row + 1, len(table.grid.rows)):
        if _is_total_row(table, row_index):
            raw_total = table.grid.value(row_index, table.amount_column)
            _require_literal(
                table.grid,
                row_index,
                table.amount_column,
                "formula_in_periodised_total",
            )
            if raw_total is not None and normalize_text(raw_total):
                totals.append(
                    require_cents(
                        raw_total,
                        location=_cell_location(
                            table.grid, row_index, table.amount_column
                        ),
                    )
                )
            continue

        raw_account = table.grid.value(row_index, table.account_column)
        raw_amount = table.grid.value(row_index, table.amount_column)
        if not normalize_text(raw_account) and not normalize_text(raw_amount):
            continue
        normalized_account = normalize_account(raw_account)
        if normalized_account is None:
            if normalize_text(raw_account) in ACCOUNT_HEADERS:
                raise Blocked("repeated_periodised_header")
            if normalize_text(raw_amount):
                raise Blocked("unrecognized_periodised_account_value")
            continue
        _require_literal(
            table.grid,
            row_index,
            table.account_column,
            "formula_in_periodised_account",
        )
        if normalized_account != account:
            raise Blocked("mixed_accounts_in_periodised_report")
        _require_literal(
            table.grid,
            row_index,
            table.amount_column,
            "formula_in_periodised_amount",
        )
        amounts.append(
            require_cents(
                raw_amount,
                location=_cell_location(table.grid, row_index, table.amount_column),
            )
        )

    if not amounts and not exact_parameter:
        raise Blocked("periodised_account_identity_not_found")
    if len(totals) > 1:
        raise Blocked("multiple_periodised_totals")
    calculated = sum(amounts, Decimal("0.00"))
    displayed_total = totals[0] if totals else None
    if displayed_total is not None and displayed_total != calculated:
        raise Blocked("periodised_total_does_not_match_details")
    return calculated, displayed_total


def _format_money(value: Decimal | None) -> str | None:
    if value is None:
        return None
    if value == Decimal("-0.00"):
        value = Decimal("0.00")
    return format(value, ".2f")


def _blocked_result(account: str, reason: str, *, location: str | None = None) -> dict[str, Any]:
    result: dict[str, Any] = {
        "account": account,
        "difference": None,
        "periodised_sum": None,
        "raw_balance_value": None,
        "status": "BLOCKED",
        "reason": reason,
    }
    if location:
        result["location"] = location
    return result


def reconcile_grids(
    raw_grids: Sequence[Grid],
    periodised_grids: dict[str, Sequence[Grid]],
    *,
    company: str,
    period: str,
    raw_filename: str = "raw.xlsx",
    periodised_filenames: dict[str, str] | None = None,
    periodised_load_errors: dict[str, ReconciliationError] | None = None,
) -> dict[str, Any]:
    filenames = periodised_filenames or {
        account: f"periodised-{account}.xlsx" for account in SUPPORTED_ACCOUNTS
    }
    load_errors = periodised_load_errors or {}
    results: list[dict[str, Any]] = []

    try:
        validate_identity(raw_grids, company, period)
        raw_table = find_table(raw_grids, RAW_AMOUNT_HEADERS, report_kind="raw_balance")
        raw_global_error: Blocked | None = None
    except Blocked as error:
        raw_table = None
        raw_global_error = error

    for account in SUPPORTED_ACCOUNTS:
        base = {
            "account": account,
            "periodised_file": filenames.get(account),
            "raw_balance_file": raw_filename,
        }
        if raw_global_error is not None:
            result = _blocked_result(
                account,
                raw_global_error.reason,
                location=raw_global_error.location,
            )
            result.update(base)
            results.append(result)
            continue

        if account in load_errors:
            error = load_errors[account]
            result = _blocked_result(account, error.reason, location=error.location)
            result.update(base)
            results.append(result)
            continue

        try:
            assert raw_table is not None
            raw_value = extract_raw_value(raw_table, account)
            account_grids = periodised_grids[account]
            validate_identity(account_grids, company, period)
            periodised_table = find_table(
                account_grids,
                REMAINING_AMOUNT_HEADERS,
                report_kind="periodised",
            )
            periodised_sum, displayed_total = extract_periodised_sum(
                periodised_table, account
            )
            difference = periodised_sum - raw_value
            status = "MATCH" if difference == Decimal("0.00") else "AVVIK"
            result = {
                **base,
                "difference": _format_money(difference),
                "displayed_periodised_total": _format_money(displayed_total),
                "periodised_sum": _format_money(periodised_sum),
                "raw_balance_value": _format_money(raw_value),
                "status": status,
            }
        except KeyError:
            result = _blocked_result(account, "periodised_workbook_missing")
            result.update(base)
        except Blocked as error:
            result = _blocked_result(account, error.reason, location=error.location)
            result.update(base)
        results.append(result)

    if any(result["status"] == "BLOCKED" for result in results):
        overall = "BLOCKED"
    elif any(result["status"] == "AVVIK" for result in results):
        overall = "AVVIK"
    else:
        overall = "MATCH"
    return {
        "accounts": results,
        "company": company,
        "overall_status": overall,
        "period": period,
    }


def load_xlsx(path: Path) -> tuple[Grid, ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise Blocked("openpyxl_dependency_missing") from error

    if path.suffix.casefold() != ".xlsx":
        raise InputError("only_xlsx_workbooks_are_supported")
    if not path.is_file():
        raise InputError("workbook_not_found")

    formula_book = None
    value_book = None
    try:
        formula_book = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
        value_book = load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
    except Exception as error:
        if formula_book is not None:
            formula_book.close()
        raise Blocked("workbook_could_not_be_opened") from error

    grids: list[Grid] = []
    try:
        if formula_book.sheetnames != value_book.sheetnames:
            raise Blocked("workbook_sheet_mismatch")
        if len(formula_book.sheetnames) > 100:
            raise Blocked("workbook_has_too_many_sheets")
        for sheet_name in formula_book.sheetnames:
            formula_sheet = formula_book[sheet_name]
            value_sheet = value_book[sheet_name]
            max_row = formula_sheet.max_row or 0
            max_column = formula_sheet.max_column or 0
            if max_row > 100_000 or max_column > 500 or max_row * max_column > 5_000_000:
                raise Blocked("workbook_dimensions_are_too_large")
            rows: list[tuple[Any, ...]] = []
            formulas: set[tuple[int, int]] = set()
            formula_rows = formula_sheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            value_rows = value_sheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
            )
            for row_index, (formula_row, value_row) in enumerate(
                zip(formula_rows, value_rows)
            ):
                values: list[Any] = []
                for column_index, (formula_cell, value_cell) in enumerate(
                    zip(formula_row, value_row)
                ):
                    values.append(value_cell.value)
                    if formula_cell.data_type == "f" or (
                        isinstance(formula_cell.value, str)
                        and formula_cell.value.startswith("=")
                    ):
                        formulas.add((row_index, column_index))
                rows.append(tuple(values))
            grids.append(Grid(sheet_name, tuple(rows), frozenset(formulas)))
    finally:
        formula_book.close()
        value_book.close()
    return tuple(grids)


def _parse_periodised_arguments(values: Iterable[str]) -> dict[str, Path]:
    parsed: dict[str, Path] = {}
    for value in values:
        account, separator, raw_path = value.partition("=")
        if not separator or not raw_path:
            raise InputError("periodised_argument_must_be_ACCOUNT_PATH")
        if account not in SUPPORTED_ACCOUNTS:
            raise InputError("unsupported_periodised_account")
        if account in parsed:
            raise InputError("duplicate_periodised_account")
        parsed[account] = Path(raw_path).expanduser()
    if set(parsed) != set(SUPPORTED_ACCOUNTS):
        raise InputError("both_supported_periodised_accounts_are_required")
    return parsed


def _build_parser() -> JsonArgumentParser:
    parser = JsonArgumentParser(description=__doc__)
    parser.add_argument("--raw-balance", required=True)
    parser.add_argument(
        "--periodised",
        required=True,
        action="append",
        help="Account and workbook path as ACCOUNT=PATH; repeat for 1510 and 1560",
    )
    parser.add_argument("--company", required=True)
    parser.add_argument("--period", required=True)
    return parser


def _emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")))


def main(argv: Sequence[str] | None = None) -> int:
    try:
        arguments = _build_parser().parse_args(argv)
        company = str(arguments.company).strip()
        period = str(arguments.period).strip()
        if not re.fullmatch(r"\d{2,10}", company):
            raise InputError("company_must_be_a_numeric_code")
        if not re.fullmatch(r"\d{6}", period) or not 1 <= int(period[4:]) <= 12:
            raise InputError("period_must_be_valid_YYYYMM")

        raw_path = Path(arguments.raw_balance).expanduser()
        periodised_paths = _parse_periodised_arguments(arguments.periodised)
        raw_grids = load_xlsx(raw_path)
        periodised_grids: dict[str, tuple[Grid, ...]] = {}
        periodised_load_errors: dict[str, ReconciliationError] = {}
        for account, path in periodised_paths.items():
            try:
                periodised_grids[account] = load_xlsx(path)
            except ReconciliationError as error:
                periodised_load_errors[account] = error
        payload = reconcile_grids(
            raw_grids,
            periodised_grids,
            company=company,
            period=period,
            raw_filename=raw_path.name,
            periodised_filenames={
                account: path.name for account, path in periodised_paths.items()
            },
            periodised_load_errors=periodised_load_errors,
        )
        _emit(payload)
        if payload["overall_status"] == "MATCH":
            return 0
        if payload["overall_status"] == "AVVIK":
            return 2
        return 3
    except InputError as error:
        _emit({"error": error.reason, "overall_status": "INPUT_ERROR"})
        return 1
    except Blocked as error:
        payload: dict[str, Any] = {
            "error": error.reason,
            "overall_status": "BLOCKED",
        }
        if error.location:
            payload["location"] = error.location
        _emit(payload)
        return 3
    except Exception:
        _emit({"error": "unexpected_internal_error", "overall_status": "BLOCKED"})
        return 3


if __name__ == "__main__":
    sys.exit(main())
